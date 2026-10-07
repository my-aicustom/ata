"""Local document extraction with page-aware PDF output and explicit scan detection."""
from __future__ import annotations
import io
from typing import Dict, Iterable


def _nonempty(values: Iterable[str]):
    return [str(v).strip() for v in values if str(v or '').strip()]


def _docx_text(doc) -> str:
    parts=[]
    parts.extend(_nonempty(p.text for p in doc.paragraphs))
    for table in doc.tables:
        for row in table.rows:
            cells=_nonempty(c.text for c in row.cells)
            if cells: parts.append(' | '.join(cells))
    seen=set()
    for section in doc.sections:
        for container,label in ((section.header,'HEADER'),(section.first_page_header,'HEADER'),(section.even_page_header,'HEADER'),
                                (section.footer,'FOOTER'),(section.first_page_footer,'FOOTER'),(section.even_page_footer,'FOOTER')):
            key=container._element
            if key in seen: continue
            seen.add(key)
            text=_nonempty(p.text for p in container.paragraphs)
            for table in container.tables:
                for row in table.rows:
                    cells=_nonempty(c.text for c in row.cells)
                    if cells: text.append(' | '.join(cells))
            if text: parts.append(f'[{label}]\n'+'\n'.join(text))
    return '\n'.join(parts).strip()


def parse_bytes(filename:str,data:bytes)->Dict:
    name=(filename or 'file.bin'); ext='.'+name.rsplit('.',1)[-1].lower() if '.' in name else ''
    if ext in {'.txt','.md','.csv','.json','.yaml','.yml'}:
        return {'type':'text','text':data.decode('utf-8',errors='replace')}
    if ext=='.pdf':
        try:
            from pypdf import PdfReader
            reader=PdfReader(io.BytesIO(data)); pages=[]; ocr_pages=[]
            for i,p in enumerate(reader.pages,1):
                text=(p.extract_text() or '').strip()
                if len(text)<24: ocr_pages.append(i)
                pages.append(f'[PAGE {i}]\n{text}')
            return {
                'type':'document','text':'\n'.join(pages).strip(),'pages':len(reader.pages),
                'needs_ocr':bool(ocr_pages),'ocr_pages':ocr_pages,
                'warnings':([f'{len(ocr_pages)} halaman minim/tanpa text layer; OCR/vision diperlukan.'] if ocr_pages else []),
            }
        except Exception as e: return {'type':'error','error':f'PDF parser gagal: {e}'}
    if ext=='.docx':
        try:
            from docx import Document
            doc=Document(io.BytesIO(data)); return {'type':'document','text':_docx_text(doc),'needs_ocr':False,'ocr_pages':[],'warnings':[]}
        except Exception as e: return {'type':'error','error':f'DOCX parser gagal: {e}'}
    if ext in {'.png','.jpg','.jpeg','.webp'}: return {'type':'image','text':f'[GAMBAR: {name}] Visual perlu dianalisis oleh model multimodal; backend tidak mengarang deskripsi.'}
    if ext in {'.mp3','.wav','.m4a','.ogg','.aac','.flac','.webm','.mp4'}: return {'type':'audio','text':f'[AUDIO: {name}] Transkripsi memerlukan provider audio yang dikonfigurasi.'}
    return {'type':'binary','text':f'[BERKAS: {name}] Format belum didukung parser lokal.'}


def render_pdf_pages(data: bytes, page_numbers, max_pages: int=8, dpi: int=144):
    """Render selected 1-based PDF pages to PNG for OCR/vision. PyMuPDF is optional."""
    try:
        import fitz
    except Exception as e:
        raise RuntimeError('PyMuPDF belum terpasang untuk OCR PDF') from e
    wanted=[]
    for n in page_numbers or []:
        try: n=int(n)
        except (TypeError,ValueError): continue
        if n>0 and n not in wanted: wanted.append(n)
        if len(wanted)>=max(1,int(max_pages)): break
    doc=fitz.open(stream=data,filetype='pdf'); out={}
    try:
        scale=max(1.0,float(dpi)/72.0); matrix=fitz.Matrix(scale,scale)
        for n in wanted:
            if n>len(doc): continue
            pix=doc[n-1].get_pixmap(matrix=matrix,alpha=False)
            out[n]=pix.tobytes('png')
    finally:
        doc.close()
    return out


def merge_pdf_page_text(source_text: str, page_texts: Dict[int,str]) -> str:
    """Replace page bodies inside `[PAGE N]` blocks while keeping page locators stable."""
    import re
    text=str(source_text or '')
    matches=list(re.finditer(r'(?im)^\[PAGE\s+(\d+)\]\s*$',text))
    if not matches: return text
    chunks=[]
    for i,m in enumerate(matches):
        n=int(m.group(1)); end=matches[i+1].start() if i+1<len(matches) else len(text)
        body=text[m.end():end].strip()
        if n in page_texts and str(page_texts[n]).strip(): body=str(page_texts[n]).strip()
        chunks.append(f'[PAGE {n}]\n{body}')
    return '\n'.join(chunks).strip()
