"""Generate a thesis DOCX from latest ATA artifacts, optionally on a campus template."""
from __future__ import annotations
import io
import re
from typing import Any, Dict, Iterable, Iterator, Optional
from docx import Document
from docx.document import Document as _Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.section import _Header, _Footer
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
from docx.shared import Pt
from .bibliography import referenced_bibliography, reference_text

_ROMAN={'i':1,'ii':2,'iii':3,'iv':4,'v':5,'vi':6,'vii':7,'viii':8,'ix':9,'x':10}

def _order(a: Dict[str,Any]):
    title=str(a.get('title') or '').strip(); m=re.match(r'(?i)^bab\s+([ivx]+|\d+)',title)
    if m:
        token=m.group(1).lower(); num=int(token) if token.isdigit() else _ROMAN.get(token,99)
        return (0,num,title.lower())
    return (1,999,title.lower())


def _iter_container_paragraphs(parent) -> Iterator[Paragraph]:
    """Yield paragraphs from a document/header/footer/cell, including nested tables."""
    for p in getattr(parent,'paragraphs',[]):
        yield p
    for table in getattr(parent,'tables',[]):
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_container_paragraphs(cell)


def _all_template_paragraphs(doc: _Document) -> Iterator[Paragraph]:
    seen=set()
    for p in _iter_container_paragraphs(doc):
        if p._p not in seen:
            seen.add(p._p); yield p
    for section in doc.sections:
        for part in (section.header, section.first_page_header, section.even_page_header,
                     section.footer, section.first_page_footer, section.even_page_footer):
            for p in _iter_container_paragraphs(part):
                if p._p not in seen:
                    seen.add(p._p); yield p


def _replace_token_in_runs(paragraph: Paragraph, token: str, value: str) -> None:
    """Replace placeholders without assigning paragraph.text, preserving run formatting.

    If a token spans runs, replacement inherits the formatting of the run containing the
    token's first character; surrounding run formatting is left intact.
    """
    if token not in paragraph.text:
        return
    while token in ''.join(r.text for r in paragraph.runs):
        runs=paragraph.runs
        full=''.join(r.text for r in runs)
        start=full.find(token)
        if start < 0: break
        end=start+len(token)
        offsets=[]; cursor=0
        for i,r in enumerate(runs):
            nxt=cursor+len(r.text); offsets.append((i,cursor,nxt)); cursor=nxt
        start_info=next((x for x in offsets if x[1] <= start < x[2]),None)
        end_info=next((x for x in offsets if x[1] < end <= x[2]),None)
        if start_info is None:
            break
        if end_info is None:
            end_info=offsets[-1]
        si,s0,_=start_info; ei,e0,_=end_info
        prefix=runs[si].text[:start-s0]
        suffix=runs[ei].text[end-e0:]
        if si==ei:
            runs[si].text=prefix+value+suffix
        else:
            runs[si].text=prefix+value
            for j in range(si+1,ei): runs[j].text=''
            runs[ei].text=suffix


def _replace_placeholders(doc: _Document, project: Dict[str,Any]) -> None:
    values={
        '{TITLE}':project.get('topic',''),
        '{STUDENT}':project.get('student',''),
        '{PROGRAM}':project.get('program',''),
        '{ADVISOR}':project.get('advisor','') or '',
        '{CAMPUS}':project.get('campus','') or '',
    }
    for p in _all_template_paragraphs(doc):
        for token,value in values.items():
            _replace_token_in_runs(p,token,str(value))


def _append_content(doc: _Document, content: str) -> None:
    for block in re.split(r'\n\s*\n',content.strip()):
        block=block.strip()
        if not block: continue
        if block.startswith('### '): doc.add_heading(block[4:].strip(),level=3)
        elif block.startswith('## '): doc.add_heading(block[3:].strip(),level=2)
        elif block.startswith('# '): doc.add_heading(block[2:].strip(),level=1)
        else:
            p=doc.add_paragraph(block)
            p.paragraph_format.space_after=Pt(6)


def build_manuscript_docx(project: Dict[str,Any], artifacts: Iterable[Dict[str,Any]], template_bytes: Optional[bytes]=None, sources: Optional[Iterable[Dict[str,Any]]]=None) -> bytes:
    doc=Document(io.BytesIO(template_bytes)) if template_bytes else Document()
    if template_bytes:
        _replace_placeholders(doc,project)
    else:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run=p.add_run(str(project.get('topic') or 'Tesis')); run.bold=True; run.font.size=Pt(16)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(str(project.get('student') or 'Mahasiswa'))
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(str(project.get('program') or 'Program Magister'))
        if project.get('advisor'):
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Pembimbing: '+str(project['advisor']))
        doc.add_page_break()
    ordered=sorted(list(artifacts),key=_order)
    for idx,a in enumerate(ordered):
        doc.add_heading(str(a.get('title') or 'Bagian'),level=1)
        _append_content(doc,str(a.get('content') or ''))
        if idx < len(ordered)-1: doc.add_page_break()
    bibliography=referenced_bibliography(ordered,list(sources or []))
    if bibliography:
        doc.add_page_break(); doc.add_heading('Daftar Pustaka',level=1)
        for source in bibliography:
            p=doc.add_paragraph(reference_text(source)); p.paragraph_format.space_after=Pt(6)
    buf=io.BytesIO(); doc.save(buf); return buf.getvalue()
