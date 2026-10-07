import io
import os
import tempfile
import unittest

from docx import Document
from pypdf import PdfWriter

from core import db
from core.defense import build_defense_context, build_balanced_thesis_context
from core.evidence import resolve_draft_evidence, verify_quote_in_document
from core.export_docx import build_manuscript_docx
from core.file_parser import parse_bytes, render_pdf_pages, merge_pdf_page_text
from core.revisions import propose_revision, accept_revision


class HardeningFinalTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.old = os.environ.get('ATA_DATA_DIR')
        os.environ['ATA_DATA_DIR'] = self.tmp.name
        db.reset_for_tests(); db.init_db()
        self.pid = db.create_project('s1', {'student':'A','program':'S2','topic':'Judul Tes','method':'interview'})

    def tearDown(self):
        if self.old is None: os.environ.pop('ATA_DATA_DIR', None)
        else: os.environ['ATA_DATA_DIR'] = self.old
        self.tmp.cleanup()

    def test_accepted_revision_preserves_source_refs(self):
        art = db.save_artifact(self.pid, 'chapter', 'Bab II', 'Versi awal.', ['src-a', 'src-b'])
        task = propose_revision(self.pid, None, art['id'], 'Versi revisi.', 'rapikan')
        accepted = accept_revision(self.pid, task['id'])
        self.assertEqual(accepted['artifact']['source_refs'], ['src-a', 'src-b'])
        versions = db.list_artifact_versions(self.pid, 'chapter', 'Bab II')
        self.assertEqual(versions[0]['source_refs'], ['src-a', 'src-b'])
        self.assertEqual(versions[1]['source_refs'], ['src-a', 'src-b'])

    def test_initial_artifact_collects_inline_brief_refs(self):
        art = db.save_artifact(self.pid, 'chapter', 'Bab I', 'Klaim [Brief: src-a, hal. 2]. Lain [Brief: 10.1000/xyz, page 4].')
        self.assertEqual(art['source_refs'], ['src-a', '10.1000/xyz'])

    def test_accepted_revision_unions_new_brief_refs_with_existing_refs(self):
        art = db.save_artifact(self.pid, 'chapter', 'Bab III', 'Dasar lama [Brief: src-a, hal. 1].', ['src-a'])
        task = propose_revision(self.pid, None, art['id'], 'Dasar baru [Brief: src-a, hal. 1]. Tambahan [Brief: src-c, hal. 3].', 'tambah sumber')
        accepted = accept_revision(self.pid, task['id'])
        self.assertEqual(accepted['artifact']['source_refs'], ['src-a', 'src-c'])

    def test_docx_placeholder_replacement_preserves_run_formatting_and_structures(self):
        template = Document()
        p = template.add_paragraph()
        r1 = p.add_run('Judul: '); r1.bold = True
        r2 = p.add_run('{TITLE}'); r2.italic = True
        table = template.add_table(rows=1, cols=1)
        table.cell(0,0).text = 'Mahasiswa: {STUDENT}'
        section = template.sections[0]
        section.header.paragraphs[0].text = 'Program {PROGRAM}'
        section.footer.paragraphs[0].text = 'Pembimbing {ADVISOR}'
        buf = io.BytesIO(); template.save(buf)

        raw = build_manuscript_docx(
            {'topic':'Tes Format','student':'Nama A','program':'S2 X','advisor':'Dr. Y'},
            [],
            buf.getvalue(),
        )
        out = Document(io.BytesIO(raw))
        self.assertEqual(out.paragraphs[0].text, 'Judul: Tes Format')
        self.assertTrue(out.paragraphs[0].runs[0].bold)
        self.assertEqual(out.paragraphs[0].runs[0].text, 'Judul: ')
        self.assertTrue(out.paragraphs[0].runs[1].italic)
        self.assertEqual(out.paragraphs[0].runs[1].text, 'Tes Format')
        self.assertIn('Nama A', out.tables[0].cell(0,0).text)
        self.assertIn('S2 X', out.sections[0].header.paragraphs[0].text)
        self.assertIn('Dr. Y', out.sections[0].footer.paragraphs[0].text)

    def test_docx_placeholder_split_across_runs_is_replaced(self):
        template=Document(); p=template.add_paragraph(); a=p.add_run('{TI'); a.italic=True; b=p.add_run('TLE}'); b.bold=True
        buf=io.BytesIO(); template.save(buf)
        raw=build_manuscript_docx({'topic':'Judul Terbagi'},[],buf.getvalue())
        out=Document(io.BytesIO(raw)); self.assertEqual(out.paragraphs[0].text,'Judul Terbagi')
        self.assertTrue(out.paragraphs[0].runs[0].italic)

    def test_defense_context_retrieves_relevant_late_chapter_not_first_n_chars(self):
        artifacts = [
            {'title':'Bab I','content':('pendahuluan umum ' * 2500)},
            {'title':'Bab IV','content':'Hasil uji menunjukkan heteroskedastisitas residual tidak terdeteksi. Temuan utama konsisten.'},
            {'title':'Bab V','content':'Kesimpulan dan keterbatasan.'},
        ]
        ctx = build_defense_context(artifacts, 'Bagaimana Anda memastikan heteroskedastisitas residual?', max_chars=5000)
        self.assertIn('heteroskedastisitas residual', ctx.lower())
        self.assertIn('Bab IV', ctx)
        self.assertLessEqual(len(ctx), 5000)

    def test_balanced_thesis_context_samples_across_artifacts(self):
        artifacts = [
            {'title':'Bab I','content':'AWALSATU ' + ('x ' * 6000) + 'AKHIRSATU'},
            {'title':'Bab V','content':'AWALLIMA ' + ('y ' * 6000) + 'AKHIRLIMA'},
        ]
        ctx = build_balanced_thesis_context(artifacts, max_chars=6000)
        self.assertIn('AWALSATU', ctx)
        self.assertIn('AKHIRSATU', ctx)
        self.assertIn('AWALLIMA', ctx)
        self.assertIn('AKHIRLIMA', ctx)

    def test_manual_evidence_is_not_verified_but_document_match_is(self):
        db.save_source(self.pid, {'id':'src1','type':'journal','title':'Paper','authors':['A'],'year':2025,'verified':True,'verification_status':'verified'})
        db.save_evidence(self.pid, 'src1', 'hal. 2', 'Temuan inti.')
        manual = resolve_draft_evidence(self.pid, 'Klaim [Brief: src1, hal. 2].')
        self.assertEqual(manual['sentences'][0]['status'], 'partial')

        source_text='[PAGE 1]\nPendahuluan.\n[PAGE 2]\nTemuan inti yang mendukung klaim secara spesifik.\n[PAGE 3]\nPenutup.'
        verification = verify_quote_in_document('Temuan inti yang mendukung klaim secara spesifik.', source_text, 'hal. 2')
        self.assertTrue(verification['verified'])
        db.save_evidence(self.pid, 'src1', 'hal. 2', 'Temuan inti yang mendukung klaim secara spesifik.', verification=verification, evidence_id='verified-ev')
        verified = resolve_draft_evidence(self.pid, 'Klaim [Brief: src1, hal. 2].')
        self.assertEqual(verified['sentences'][0]['status'], 'supported')
        self.assertEqual(verified['sentences'][0]['resolved'][0]['verification_status'], 'verified')

    def test_non_doi_source_can_be_supported_by_verified_document_chain(self):
        db.save_source(self.pid, {'id':'reg1','type':'regulation','title':'Peraturan X','authors':['Instansi'],'year':2026,'verified':False,'verification_status':'candidate'})
        document='[PAGE 5]\nKetentuan wajib diterapkan pada seluruh unit kerja.'
        verification=verify_quote_in_document('Ketentuan wajib diterapkan pada seluruh unit kerja.',document,'hal. 5')
        db.save_evidence(self.pid,'reg1','hal. 5','Ketentuan wajib diterapkan pada seluruh unit kerja.',verification=verification)
        out=resolve_draft_evidence(self.pid,'Aturan berlaku [Brief: reg1, hal. 5].')
        resolved=out['sentences'][0]['resolved'][0]
        self.assertEqual(out['sentences'][0]['status'],'supported')
        self.assertFalse(resolved['source_verified'])
        self.assertEqual(resolved['support_basis'],'document_backed')
        self.assertEqual(resolved['bibliographic_status'],'candidate')

    def test_docx_parser_extracts_tables_headers_and_footers(self):
        doc = Document(); doc.add_paragraph('Isi utama')
        doc.add_table(rows=1, cols=1).cell(0,0).text='Isi tabel penting'
        doc.sections[0].header.paragraphs[0].text='Header kampus'
        doc.sections[0].footer.paragraphs[0].text='Footer halaman'
        buf=io.BytesIO(); doc.save(buf)
        parsed=parse_bytes('tesis.docx',buf.getvalue())
        self.assertEqual(parsed['type'],'document')
        self.assertIn('Isi utama',parsed['text'])
        self.assertIn('Isi tabel penting',parsed['text'])
        self.assertIn('Header kampus',parsed['text'])
        self.assertIn('Footer halaman',parsed['text'])

    def test_pdf_ocr_helpers_render_and_merge_selected_page(self):
        writer=PdfWriter(); writer.add_blank_page(width=200,height=200)
        buf=io.BytesIO(); writer.write(buf)
        rendered=render_pdf_pages(buf.getvalue(),[1],max_pages=1,dpi=72)
        self.assertTrue(rendered[1].startswith(b'\x89PNG'))
        merged=merge_pdf_page_text('[PAGE 1]\n', {1:'Teks hasil OCR'})
        self.assertIn('[PAGE 1]',merged)
        self.assertIn('Teks hasil OCR',merged)

    def test_blank_pdf_is_flagged_as_needing_ocr(self):
        writer=PdfWriter(); writer.add_blank_page(width=300,height=300)
        buf=io.BytesIO(); writer.write(buf)
        parsed=parse_bytes('scan.pdf',buf.getvalue())
        self.assertEqual(parsed['type'],'document')
        self.assertTrue(parsed['needs_ocr'])
        self.assertIn(1, parsed['ocr_pages'])


if __name__ == '__main__': unittest.main()
